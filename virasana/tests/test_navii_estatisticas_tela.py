"""Testes da aba "Estatísticas de uso" (F2b): página, API, cálculo do universo em partes e planilha — com o cenário
do cálculo (test_navii_estatisticas_uso: pesquisas de dez/2025 a mar/2026, setores e Supervisor).
"""
import io
import os
import unittest
from unittest import mock

import mongomock
from flask import Flask
from flask_login import LoginManager, UserMixin
from openpyxl import load_workbook

from virasana.navii import busca_navii, coleta_uso, tabelas_opcionais, views_navii
from virasana.navii.views_navii import bp_navii
from virasana.tests.test_navii_busca_sqlite import cria_sessao, monta_cenario
from virasana.tests.test_navii_estatisticas_uso import ANA, BETO, SUP, cria_fichas, registra_pesquisas


class Usuario(UserMixin):
    def __init__(self, id_):
        self.id = id_


class TelaTestCase(unittest.TestCase):

    def setUp(self):
        tabelas_opcionais.limpa_cache()
        coleta_uso.limpa_cache()
        self.session = cria_sessao()
        monta_cenario(self.session)
        cria_fichas(self.session)
        item = {c: i for i, c in self.session.query(busca_navii.Item.ID, busca_navii.Item.codigoConteiner)}
        self.db = mongomock.MongoClient()['virasana']
        registra_pesquisas(coleta_uso.colecao(self.db), item)
        self.app = Flask(__name__, template_folder=os.path.join(os.path.dirname(os.path.dirname(__file__)),
                                                                'templates'))
        self.app.config.update(TESTING=True, db_session=self.session, mongodb=self.db, SECRET_KEY='x')
        LoginManager(self.app).request_loader(
            lambda req: Usuario(req.headers['X-Usuario']) if req.headers.get('X-Usuario') else None)

        @self.app.route('/mini1/<_id>')
        def mini(_id):
            return ''
        self.app.register_blueprint(bp_navii)
        self.client = self.app.test_client()

    def tearDown(self):
        self.session.close()

    def api(self, usuario=ANA, **parametros):
        parametros = {'tipo': 'mes', 'codigo': '2026-01', **parametros}
        return self.client.get('/navii/api/estatisticas', query_string=parametros,
                               headers={'X-Usuario': usuario} if usuario else {})

    def universo(self, pendentes, usuario=ANA):
        return self.client.post('/navii/api/estatisticas/universo', json=pendentes,
                                headers={'X-Usuario': usuario} if usuario else {})

    def calcula(self):
        pendentes = self.api().get_json()['pendentes']
        while pendentes:
            pendentes = self.universo({'pendentes': pendentes}).get_json()['restantes']


class TestPaginaEApi(TelaTestCase):

    def test_pagina_recebe_as_urls_e_tem_as_abas(self):
        html = self.client.get('/navii/estatisticas', headers={'X-Usuario': ANA}).get_data(as_text=True)
        for trecho in ('var API_URL = "/navii/api/estatisticas"',
                       'var UNIVERSO_URL = "/navii/api/estatisticas/universo"',
                       'var PLANILHA_URL = "/navii/estatisticas/planilha"', 'aria-current="page">Estatísticas de uso',
                       'href="/navii/fichas"', 'name="csrf-token"'):
            self.assertIn(trecho, html)
        for tela in ('/navii/', '/navii/fichas'):                            # a aba nova nas outras telas
            self.assertIn('href="/navii/estatisticas">Estatísticas de uso',
                          self.client.get(tela, headers={'X-Usuario': ANA}).get_data(as_text=True))

    def test_pendentes_calculados_em_partes_depois_os_numeros(self):
        dados = self.api().get_json()
        self.assertEqual((len(dados['pendentes']), dados['cobertura']['atual']), (43, None))   # dias
        with mock.patch.object(views_navii, 'ORCAMENTO_UNIVERSO', -1):       # uma parte por chamada (sempre)
            resposta = self.universo({'pendentes': dados['pendentes']}).get_json()
        self.assertEqual((resposta['calculados'], len(resposta['restantes'])), (1, 42))
        self.calcula()
        dados = self.api().get_json()
        self.assertEqual(dados['pendentes'], [])
        self.assertEqual((dados['cobertura']['atual']['percentual'], dados['seletividade']['mediana']), (75.0, 50.0))
        self.assertEqual([linha['nome'] for linha in dados['usuarios']], ['Ana'])

    def test_erros(self):
        self.assertEqual(self.api(usuario=None).status_code, 401)
        self.assertEqual(self.universo({'pendentes': []}, usuario=None).status_code, 401)
        self.assertEqual(self.api(tipo='dia').status_code, 400)
        self.assertEqual(self.api(codigo='2026-13').status_code, 400)
        for corpo in ({'pendentes': 'x'}, ['lista'], {'pendentes': [{'janela': 'efetiva', 'inicio': '2026-01-01',
                                                                     'fim': '2026-01-31'}]}):   # um mês: não
            with self.subTest(corpo=corpo):
                self.assertEqual(self.universo(corpo).status_code, 400)
        with mock.patch.object(views_navii, 'MAX_PENDENTES', 1):
            pedaco = {'janela': 'efetiva', 'inicio': '2026-01-01', 'fim': '2026-01-01'}
            self.assertEqual(self.universo({'pendentes': [pedaco, pedaco]}).status_code, 400)
        del self.app.config['mongodb']
        self.assertEqual(self.api().status_code, 503)
        self.assertEqual(self.universo({'pendentes': []}).status_code, 503)

    def test_porto_configuravel_e_log_de_uso(self):
        self.app.config['navii_porto_estatisticas'] = 'BRRIG'
        with self.assertLogs('virasana.navii.uso', level='INFO') as log:
            dados = self.api().get_json()
        self.assertEqual(dados['porto'], 'BRRIG')
        self.assertTrue(any('acao=estatisticas' in linha and 'periodo=2026-01' in linha for linha in log.output))


class TestPlanilha(TelaTestCase):

    def planilha(self, usuario):
        resposta = self.client.get('/navii/estatisticas/planilha', query_string={'tipo': 'mes', 'codigo': '2026-01'},
                                   headers={'X-Usuario': usuario})
        self.assertEqual(resposta.status_code, 200)
        self.assertIn('estatisticas_navii_2026-01.xlsx', resposta.headers['Content-Disposition'])
        return load_workbook(io.BytesIO(resposta.data))

    def nomes(self, livro, aba):
        return [linha[0] for linha in livro[aba].iter_rows(min_row=2, values_only=True)]

    def test_abas_e_acesso(self):
        self.calcula()
        livro = self.planilha(SUP)
        self.assertEqual(livro.sheetnames, ['Resumo', 'Por setor', 'Por usuário', 'Por dia'])
        self.assertEqual(self.nomes(livro, 'Por usuário'), ['Abel', 'Ana'])
        self.assertEqual(self.nomes(self.planilha(BETO), 'Por usuário'), ['Beto'])
        self.assertEqual(self.nomes(livro, 'Por setor'), ['EQRIMP', 'SACIT', 'SUB'])
        resumo = {linha[0]: linha[1:] for linha in livro['Resumo'].iter_rows(values_only=True) if linha[0]}
        self.assertEqual(resumo['Cobertura (%)'][:2], (75.0, None))
        self.assertEqual(resumo['Pesquisas'], (3, 1, 200.0))                 # 3 contra 1: +200%
        self.assertEqual(resumo['Universo'][0], 'completo')

    def test_universo_em_calculo_e_sem_usuario(self):
        livro = self.planilha(ANA)
        resumo = {linha[0]: linha[1:] for linha in livro['Resumo'].iter_rows(values_only=True) if linha[0]}
        self.assertEqual(resumo['Universo'][0], 'em cálculo')
        self.assertEqual(self.client.get('/navii/estatisticas/planilha').status_code, 401)


if __name__ == '__main__':
    unittest.main()
